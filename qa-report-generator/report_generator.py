"""
QA Report Generator
--------------------
Turns a JSON list of test/verification results into a styled Pass/Fail
Excel report. Built to remove the manual spreadsheet upkeep that QA teams
otherwise do by hand after every test run.

Input format (see sample_results.json):
[
  {
    "name": "Checkout: add_to_cart event fires",
    "category": "Event Verification",
    "status": "pass" | "fail",
    "expected": "string describing what was expected",
    "actual": "string describing what actually happened/was observed",
    "notes": "optional free-text notes"
  },
  ...
]

Usage:
    python report_generator.py sample_results.json sample_output/report.xlsx
"""

import sys
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
PASS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
PASS_FONT = Font(color="166534", bold=True)
FAIL_FONT = Font(color="991B1B", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

COLUMNS = ["#", "Test Case", "Category", "Status", "Expected", "Actual", "Notes"]


def load_results(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def build_report(results, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    total = len(results)
    passed = sum(1 for r in results if r.get("status", "").lower() == "pass")
    failed = total - passed
    pass_rate = (passed / total * 100) if total else 0

    # Summary block
    ws["A1"] = "QA Verification Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(italic=True, size=9, color="6B7280")
    ws["A4"] = "Total"
    ws["B4"] = total
    ws["A5"] = "Passed"
    ws["B5"] = passed
    ws["B5"].font = PASS_FONT
    ws["A6"] = "Failed"
    ws["B6"] = failed
    ws["B6"].font = FAIL_FONT
    ws["A7"] = "Pass rate"
    ws["B7"] = f"{pass_rate:.1f}%"
    for row in range(4, 8):
        ws[f"A{row}"].font = Font(bold=True)

    header_row = 9
    for col_idx, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for i, result in enumerate(results, start=1):
        row = header_row + i
        status = result.get("status", "").lower()
        values = [
            i,
            result.get("name", ""),
            result.get("category", ""),
            "PASS" if status == "pass" else "FAIL",
            result.get("expected", ""),
            result.get("actual", ""),
            result.get("notes", ""),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 4:
                cell.fill = PASS_FILL if status == "pass" else FAIL_FILL
                cell.font = PASS_FONT if status == "pass" else FAIL_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [5, 38, 20, 10, 32, 32, 28]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = f"A{header_row + 1}"

    wb.save(output_path)
    return total, passed, failed


def main():
    if len(sys.argv) != 3:
        print("Usage: python report_generator.py <results.json> <output.xlsx>")
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]
    results = load_results(input_path)
    total, passed, failed = build_report(results, output_path)
    print(f"Report written to {output_path}")
    print(f"{passed}/{total} passed, {failed} failed")


if __name__ == "__main__":
    main()
